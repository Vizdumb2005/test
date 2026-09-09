from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

from docx import Document as DocxDocument

from .base import BaseLoader
from src.core.models import Document


class DocxLoader(BaseLoader):
    document_type = "docx"

    def __init__(
        self,
        metadata: Optional[dict[str, Any]] = None,
        extract_tables: bool = False,
    ) -> None:
        super().__init__(metadata=metadata)
        self.extract_tables = extract_tables

    def _load(self, source: str) -> Document:
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"DOCX file not found: {source}")

        doc = DocxDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs)
        document_id = f"docx_{path.stem}_{hash(path) & 0xFFFFFFFF:X}"

        table_count = len(doc.tables) if self.extract_tables else 0

        return Document(
            document_id=document_id,
            source=str(path.resolve()),
            file_name=path.name,
            document_type=self.document_type,
            text=text,
            metadata={
                "paragraph_count": len(paragraphs),
                "table_count": table_count,
                "extract_tables": self.extract_tables,
                "size_bytes": path.stat().st_size,
                **self.metadata,
            },
            page_count=None,
        )
